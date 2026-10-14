#!/usr/bin/env python3
"""Fase 1b — Volcado completo del Excel a JSON + localización de la anotación 127→137."""
import json
from openpyxl import load_workbook

XLSX = '/home/z/my-project/upload/Costpro.xlsx'
OUT = '/home/z/my-project/scripts/enervida-recon/excel_dump.json'

wb = load_workbook(XLSX, data_only=True, read_only=True)

def sheet_rows(name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else f'col{i}' for i, h in enumerate(rows[0])]
    out = []
    for idx, r in enumerate(rows[1:], start=2):  # fila Excel real (encabezado = 1)
        d = {hdr[i]: r[i] if i < len(r) else None for i in range(len(hdr))}
        d['_excel_row'] = idx
        out.append(d)
    return hdr, out

dump = {}
for name in wb.sheetnames:
    hdr, data = sheet_rows(name)
    # normalizar fechas a ISO
    for d in data:
        for k, v in list(d.items()):
            if hasattr(v, 'isoformat'):
                d[k] = v.isoformat()
    dump[name] = {'headers': hdr, 'rows': data}
    json.dump(dump[name], open(f'/home/z/my-project/scripts/enervida-recon/{name.lower()}.json', 'w'), ensure_ascii=False, default=str)

json.dump(dump, open(OUT, 'w'), ensure_ascii=False, default=str)

# Localizar la anotación de rectificación
op = dump['Operaciones']['rows']
annot = [d for d in op if any(isinstance(v, str) and 'rectific' in v.lower() for v in d.values())]
print('=== ANOTACIÓN 127→137 ===')
for a in annot:
    print(json.dumps(a, ensure_ascii=False, default=str))

# Fila 873 exacta (encabezado = fila 1)
row873 = [d for d in op if d['_excel_row'] == 873]
print('\n=== FILA 873 (Excel) ===')
print(json.dumps(row873, ensure_ascii=False, default=str))

# Vecindario de la fila 873 (870-876) para contexto
print('\n=== FILAS 870-876 ===')
for d in op:
    if 870 <= d['_excel_row'] <= 876:
        print(json.dumps(d, ensure_ascii=False, default=str))

# Resumen por clasificación y UM
from collections import Counter
print('\nClasificación × UM:', dict(Counter((d['Clasificación'], d['UM']) for d in op)))
skus = sorted({d['Código'] for d in op})
print(f'SKUs distintos en Operaciones: {len(skus)} (min={min(skus)}, max={max(skus)})')
nf = [d for d in op if isinstance(d['Cantidad'], float)]
print(f'Filas con cantidad decimal: {len(nf)}; muestras: {[(d["_excel_row"], d["Código"], d["Cantidad"]) for d in nf[:8]]}')
exist = dump['ExistenciaFinal']['rows']
comi = dump['ConfComisiones']['rows']
print(f'SKUs ExistenciaFinal: {len(exist)} (min={min(d["Código"] for d in exist)}, max={max(d["Código"] for d in exist)})')
print(f'SKUs ConfComisiones: {len(comi)}; comisión>0: {sum(1 for d in comi if (d["Comisión por Unidad Vendida"] or 0) > 0)}; =0: {sum(1 for d in comi if d["Comisión por Unidad Vendida"] == 0)}; vacía: {sum(1 for d in comi if d["Comisión por Unidad Vendida"] is None)}')
print(f'\nDump completo: {OUT}')
