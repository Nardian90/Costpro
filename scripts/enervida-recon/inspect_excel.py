#!/usr/bin/env python3
"""Fase 1 — Inspección integral de Costpro.xlsx (solo lectura)."""
import json
from openpyxl import load_workbook

XLSX = '/home/z/my-project/upload/Costpro.xlsx'

wb = load_workbook(XLSX, data_only=True, read_only=True)
print('=== HOJAS ===')
print(wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    print(f'\n{"="*70}\n### HOJA: {name}  (max_row={ws.max_row}, max_col={ws.max_column})')
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print('  (vacía)')
        continue
    headers = [str(h) if h is not None else f'col{i}' for i, h in enumerate(rows[0])]
    print(f'ENCABEZADOS ({len(headers)}): {headers}')
    data = rows[1:]
    print(f'REGISTROS: {len(data)}')
    for i, h in enumerate(headers):
        types = {}
        empties = 0
        samples = []
        for r in data:
            v = r[i] if i < len(r) else None
            if v is None or (isinstance(v, str) and v.strip() == ''):
                empties += 1
                continue
            t = type(v).__name__
            types[t] = types.get(t, 0) + 1
            if len(samples) < 3 and repr(v) not in [repr(s) for s in samples]:
                samples.append(v)
        print(f'  {h:<15} tipos={types} vacíos={empties} muestras={samples[:3]}')
    print('  PRIMERAS FILAS:')
    for r in data[:5]:
        print('   ', r)

if 'Operaciones' in wb.sheetnames:
    ws = wb['Operaciones']
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) if h is not None else f'col{i}' for i, h in enumerate(rows[0])]
    ci = {h: i for i, h in enumerate(hdr)}
    data = rows[1:]
    cls = {}
    dates = []
    for r in data:
        c = r[ci['Clasificación']] if 'Clasificación' in ci else None
        cls[str(c)] = cls.get(str(c), 0) + 1
        f = r[ci['Fecha']] if 'Fecha' in ci else None
        if f is not None:
            dates.append(f)
    print(f'\n### Operaciones — Clasificación: {cls}')
    if dates:
        try:
            print(f'    Fechas: min={min(dates)} max={max(dates)}')
        except Exception:
            print(f'    Fechas (sin ordenar): muestras={dates[:5]} tipos={ {type(d).__name__ for d in dates} }')
